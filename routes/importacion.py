from __future__ import annotations
import csv
import io
from flask import Blueprint, render_template, request, redirect, url_for, flash
from routes.helpers import roles_permitidos, auditoria
from database.db import consultar, ejecutar
from roles import es_correo_gmail

importacion_bp = Blueprint("importacion", __name__, url_prefix="/importacion")
ROLES = ("Super Admin", "Rector", "Administrador / Coordinador")


def _leer_archivo(file):
    nombre = (file.filename or "").lower()
    if nombre.endswith(".csv"):
        raw = file.read().decode("utf-8-sig")
        return list(csv.DictReader(io.StringIO(raw)))
    if nombre.endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(file, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(x or "").strip() for x in rows[0]]
        return [dict(zip(headers, row)) for row in rows[1:]]
    raise ValueError("Solo se admiten archivos .csv o .xlsx.")


def _get(row, *keys):
    for key in keys:
        if key in row and str(row[key] or "").strip():
            return row[key]
    return ""


def _lookup(query, value):
    return consultar(query, (str(value or "").strip(),), uno=True)


@importacion_bp.route("/")
@roles_permitidos(*ROLES)
def index():
    return render_template("importacion/index.html")


@importacion_bp.post("/estudiantes")
@roles_permitidos(*ROLES)
def estudiantes():
    file = request.files.get("archivo")
    if not file or not file.filename:
        flash("Selecciona un archivo CSV o XLSX.", "error")
        return redirect(url_for("importacion.index"))
    try:
        rows = _leer_archivo(file)
    except Exception as exc:
        flash(str(exc), "error")
        return redirect(url_for("importacion.index"))

    ok = 0
    errores = []
    for line, row in enumerate(rows, start=2):
        try:
            codigo = str(_get(row, "codigo", "codigo_estudiantil")).strip()
            tipo = str(_get(row, "tipo_documento") or "TI").strip()
            documento = str(_get(row, "documento")).strip()
            nombres = str(_get(row, "nombres")).strip()
            apellidos = str(_get(row, "apellidos")).strip()
            sede = _lookup("SELECT id FROM sedes WHERE estado=1 AND LOWER(nombre)=LOWER(%s) LIMIT 1", _get(row, "sede", "sede_nombre"))
            jornada = _lookup("SELECT id FROM jornadas WHERE estado=1 AND LOWER(nombre)=LOWER(%s) LIMIT 1", _get(row, "jornada", "jornada_nombre"))
            periodo = _lookup("SELECT id FROM periodos_academicos WHERE nombre=%s LIMIT 1", _get(row, "periodo", "periodo_nombre"))
            if not all((codigo, documento, nombres, apellidos, sede, jornada, periodo)):
                raise ValueError("Faltan datos obligatorios o una referencia institucional no existe.")
            if consultar("SELECT id_estudiante FROM estudiantes WHERE documento=%s OR codigo_estudiantil=%s", (documento, codigo), uno=True):
                raise ValueError("Documento o código estudiantil duplicado.")
            grado = _lookup("SELECT id FROM grados WHERE nombre=%s LIMIT 1", _get(row, "grado")) if _get(row, "grado") else None
            grupo = consultar("SELECT id FROM grupos WHERE nombre=%s AND id_grado=%s LIMIT 1", (str(_get(row, "grupo")).strip(), grado["id"]), uno=True) if grado and _get(row, "grupo") else None
            ejecutar(
                """INSERT INTO estudiantes(codigo_estudiantil,tipo_documento,documento,nombres,apellidos,curso,id_sede,id_jornada,id_periodo,id_grado,id_grupo,estado)
                   VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,'Activo')""",
                (codigo, tipo, documento, nombres, apellidos, _get(row, "curso"), sede["id"], jornada["id"], periodo["id"], grado["id"] if grado else None, grupo["id"] if grupo else None),
            )
            ok += 1
        except Exception as exc:
            errores.append(f"Fila {line}: {exc}")
    auditoria("IMPORTAR", "Estudiantes", None, f"Correctos: {ok}; errores: {len(errores)}")
    flash(f"Importación finalizada: {ok} creados y {len(errores)} con error.", "success" if not errores else "error")
    for item in errores[:10]:
        flash(item, "error")
    return redirect(url_for("importacion.index"))


@importacion_bp.post("/acudientes")
@roles_permitidos(*ROLES)
def acudientes():
    file = request.files.get("archivo")
    if not file or not file.filename:
        flash("Selecciona un archivo CSV o XLSX.", "error")
        return redirect(url_for("importacion.index"))
    try:
        rows = _leer_archivo(file)
    except Exception as exc:
        flash(str(exc), "error")
        return redirect(url_for("importacion.index"))

    ok = 0
    errores = []
    for line, row in enumerate(rows, start=2):
        try:
            tipo = str(_get(row, "tipo_documento") or "CC").strip()
            documento = str(_get(row, "documento")).strip()
            nombres = str(_get(row, "nombres")).strip()
            apellidos = str(_get(row, "apellidos")).strip()
            correo = str(_get(row, "correo")).strip().lower()
            if not all((documento, nombres, apellidos)):
                raise ValueError("Faltan documento, nombres o apellidos.")
            if not es_correo_gmail(correo):
                raise ValueError("El correo del acudiente debe terminar en @gmail.com.")
            if consultar("SELECT id_acudiente FROM acudientes WHERE documento=%s", (documento,), uno=True):
                raise ValueError("Documento duplicado.")
            ejecutar(
                """INSERT INTO acudientes(tipo_documento,documento,nombres,apellidos,telefono,correo,direccion,barrio,municipio,estado)
                   VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,1)""",
                (tipo, documento, nombres, apellidos, _get(row, "telefono"), correo, _get(row, "direccion"), _get(row, "barrio"), _get(row, "municipio") or "Ibagué"),
            )
            ok += 1
        except Exception as exc:
            errores.append(f"Fila {line}: {exc}")
    auditoria("IMPORTAR", "Acudientes", None, f"Correctos: {ok}; errores: {len(errores)}")
    flash(f"Importación finalizada: {ok} creados y {len(errores)} con error.", "success" if not errores else "error")
    for item in errores[:10]:
        flash(item, "error")
    return redirect(url_for("importacion.index"))
